def is_valid_text(content):
    if not content or not content.strip():
        return False

    length = len(content)

    non_text_ratio = sum(
        1 for c in content
        if not c.isalnum() and not c.isspace()
    ) / max(length, 1)

    alpha_ratio = sum(
        1 for c in content
        if c.isalpha()
    ) / max(length, 1)

    if non_text_ratio > 0.4:
        return False

    if alpha_ratio < 0.3:
        return False

    return True


def read_pdf(file_path):
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        text = ""

        for page in reader.pages:
            text += page.extract_text() or ""

        return text
    except Exception as e:
        return f"Error reading PDF: {e}"


def read_text(file_path):
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except Exception as e:
        return f"Error reading file: {e}"


def read_docx(file_path):
    try:
        from docx import Document

        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        return f"Error reading DOCX: {e}"


def read_pptx(file_path):
    try:
        from pptx import Presentation

        prs = Presentation(file_path)
        text = []

        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)

        return "\n".join(text)
    except Exception as e:
        return f"Error reading PPTX: {e}"


def read_excel(file_path):
    try:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, data_only=True)
        text = []

        for sheet in wb.worksheets:
            text.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) for cell in row if cell is not None])
                if row_text:
                    text.append(row_text)

        return "\n".join(text)
    except Exception as e:
        return f"Error reading Excel: {e}"
